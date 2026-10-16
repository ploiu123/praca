"""Warsztat 2 - korelacja i DAG.

Dodaje do pliku z danymi arkusz "Etap I" z korelacjami Pearsona,
linią regresji i wykresem punktowym.

Uruchomienie:  python3 etap1.py <plik_z_danymi.xlsx> <plik_wynikowy.xlsx>
"""
import sys

from openpyxl import load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.trendline import Trendline, TrendlineLabel
from openpyxl.styles import Alignment, Border, Font, Side

WEJSCIE, WYJSCIE = sys.argv[1], sys.argv[2]
D = "'Dane syntetyczne'!"
CIENKA = Side(style="thin", color="000000")
RAMKA = Border(left=CIENKA, right=CIENKA, top=CIENKA, bottom=CIENKA)


def kol(litera, ost):
    return f"{D}${litera}$2:${litera}${ost}"


def tabela(ws, r, tytul, naglowki, wiersze, fmt):
    ws.cell(row=r, column=1, value=tytul).font = Font(bold=True, size=12)
    r += 1
    for c, n in enumerate(naglowki, 1):
        cell = ws.cell(row=r, column=c, value=n)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = RAMKA
    for nazwa, *formuly in wiersze:
        r += 1
        ws.cell(row=r, column=1, value=nazwa).border = RAMKA
        for c, f in enumerate(formuly, 2):
            cell = ws.cell(row=r, column=c, value=f)
            cell.number_format = fmt
            cell.border = RAMKA
    return r + 2


wb = load_workbook(WEJSCIE)
ws = wb.create_sheet("Etap I")

r = tabela(ws, 1, "Współczynnik korelacji Pearsona",
           ["Para zmiennych", "Wszystkie miesiące", "Bez listopada i grudnia"],
           [(nazwa, f"=PEARSON({kol(x, 13)},{kol(y, 13)})", f"=PEARSON({kol(x, 11)},{kol(y, 11)})")
            for nazwa, x, y in (("Reklama – sprzedaż", "B", "C"),
                                ("Popyt – sprzedaż", "D", "C"),
                                ("Popyt – reklama", "D", "B"))],
           "0.000")

r = tabela(ws, r, "Linia regresji: sprzedaż = a × reklama + b",
           ["", "Wszystkie miesiące", "Bez listopada i grudnia"],
           [("Nachylenie a", f"=SLOPE({kol('C', 13)},{kol('B', 13)})", f"=SLOPE({kol('C', 11)},{kol('B', 11)})"),
            ("Wyraz wolny b", f"=INTERCEPT({kol('C', 13)},{kol('B', 13)})", f"=INTERCEPT({kol('C', 11)},{kol('B', 11)})"),
            ("R²", f"=RSQ({kol('C', 13)},{kol('B', 13)})", f"=RSQ({kol('C', 11)},{kol('B', 11)})")],
           "0.000")

ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 16

dane = wb["Dane syntetyczne"]
wykres = ScatterChart()
wykres.title = "Wydatki na reklamę a sprzedaż"
wykres.style = 2
wykres.scatterStyle = "marker"
wykres.x_axis.title = "Reklama (tys. zł)"
wykres.y_axis.title = "Sprzedaż (tys. zł)"
wykres.x_axis.delete = False
wykres.y_axis.delete = False
wykres.legend = None
seria = Series(Reference(dane, min_col=3, min_row=2, max_row=13),
               Reference(dane, min_col=2, min_row=2, max_row=13), title="Miesiące")
seria.marker.symbol = "circle"
seria.marker.size = 7
seria.marker.graphicalProperties.solidFill = "4472C4"
seria.marker.graphicalProperties.line.solidFill = "4472C4"
seria.graphicalProperties.line.noFill = True
seria.trendline = Trendline(
    trendlineType="linear", dispEq=True, dispRSqr=True,
    trendlineLbl=TrendlineLabel(numFmt=NumFmt(formatCode="0.00", sourceLinked=False)))
wykres.series.append(seria)
wykres.width, wykres.height = 16, 9
ws.add_chart(wykres, "E2")

wb.save(WYJSCIE)
print("Zapisano", WYJSCIE)
