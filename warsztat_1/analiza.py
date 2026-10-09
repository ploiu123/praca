"""Warsztat 1 - filtrowanie danych sklepu internetowego.

Wczytuje dane.csv, wypisuje zestawienia dla każdego filtra z zadania
i buduje skoroszyt warsztat_1.xlsx (arkusz z filtrami + podsumowania na formułach).

Uruchomienie:  python3 analiza.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

KATALOG = Path(__file__).resolve().parent
DANE = KATALOG / "dane.csv"
XLSX = KATALOG / "warsztat_1.xlsx"

FILTRY = [
    ("zrodlo_ruchu", "Źródło ruchu"),
    ("urzadzenie", "Urządzenie"),
    ("kraj", "Kraj"),
    ("typ_uzytkownika", "Nowy / Powracający"),
]


def zestawienie(df, grupy):
    t = df.groupby(grupy).agg(
        uzytkownicy=("zakup", "size"),
        zakupy=("kupil", "sum"),
        sr_odrzucen_proc=("wskaznik_odrzucen_proc", "mean"),
        sr_czas_min=("sredni_czas_min", "mean"),
        sr_stron=("liczba_stron", "mean"),
        przychod_eur=("wartosc_koszyka_eur", "sum"),
    )
    t["konwersja_proc"] = t.zakupy / t.uzytkownicy * 100
    t["sr_koszyk_eur"] = t.przychod_eur / t.zakupy.where(t.zakupy > 0)
    t["przychod_na_uzytk_eur"] = t.przychod_eur / t.uzytkownicy
    return t.round(2)


def wypisz(df):
    pd.set_option("display.width", 200)
    for kol, nazwa in FILTRY:
        print(f"\n=== {nazwa} ===")
        print(zestawienie(df, kol).to_string())
    for kol, nazwa in FILTRY[:3]:
        print(f"\n=== {nazwa} x Nowy/Powracający ===")
        print(zestawienie(df, [kol, "typ_uzytkownika"]).to_string())
    print("\n=== Korelacja wskaźników z zakupem ===")
    print(df[["wskaznik_odrzucen_proc", "sredni_czas_min", "liczba_stron",
              "wartosc_koszyka_eur", "kupil"]].corr().round(2).to_string())


# ---------------------------------------------------------------- Excel

FONT = "Arial"
NAGLOWEK = PatternFill("solid", fgColor="1F4E78")
SEKCJA = PatternFill("solid", fgColor="DDEBF7")
CIENKA = Side(style="thin", color="BFBFBF")
RAMKA = Border(left=CIENKA, right=CIENKA, top=CIENKA, bottom=CIENKA)

KOLUMNY_DANYCH = [
    ("Użytkownik", "uzytkownik", "0"),
    ("Źródło ruchu", "zrodlo_ruchu", None),
    ("Urządzenie", "urzadzenie", None),
    ("Kraj", "kraj", None),
    ("Nowy/Powracający", "typ_uzytkownika", None),
    ("Wskaźnik odrzuceń", "wskaznik_odrzucen_proc", "0%"),
    ("Średni czas na stronie (min)", "sredni_czas_min", "0.0"),
    ("Liczba przeglądanych stron", "liczba_stron", "0"),
    ("Wartość koszyka (€)", "wartosc_koszyka_eur", "#,##0 €"),
    ("Czy dokonano zakupu?", "zakup", None),
]


def zakres(litera, n):
    return f"Dane!${litera}$2:${litera}${n + 1}"


def arkusz_dane(wb, df):
    ws = wb.active
    ws.title = "Dane"
    for c, (nagl, _, _) in enumerate(KOLUMNY_DANYCH, 1):
        ws.cell(row=1, column=c, value=nagl)
    for r, wiersz in enumerate(df.itertuples(index=False), 2):
        for c, (_, kol, fmt) in enumerate(KOLUMNY_DANYCH, 1):
            v = getattr(wiersz, kol)
            if kol == "wskaznik_odrzucen_proc":
                v = v / 100  # procent przechowywany jako ułamek
            cell = ws.cell(row=r, column=c, value=v.item() if hasattr(v, "item") else v)
            cell.font = Font(name=FONT, color="0000FF")  # dane wejściowe
            if fmt:
                cell.number_format = fmt
    ost = get_column_letter(len(KOLUMNY_DANYCH))
    tabela = Table(displayName="TabelaDane", ref=f"A1:{ost}{len(df) + 1}")
    tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabela)
    for c in range(1, len(KOLUMNY_DANYCH) + 1):
        ws.cell(row=1, column=c).font = Font(name=FONT, bold=True, color="FFFFFF")
        ws.cell(row=1, column=c).alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "A2"
    uw = len(df) + 3
    ws.cell(row=uw, column=1, value=(
        "Źródło: dane przykładowe z karty warsztatu (Ćwiczenie I). "
        "Filtry: strzałki w nagłówkach tabeli. Niebieski tekst = dane wejściowe.")
    ).font = Font(name=FONT, italic=True, color="595959")


METRYKI = [
    ("Użytkownicy", "0"),
    ("Zakupy", "0"),
    ("Współczynnik konwersji", "0.0%"),
    ("Śr. wskaźnik odrzuceń", "0.0%"),
    ("Śr. czas na stronie (min)", "0.00"),
    ("Śr. liczba stron", "0.0"),
    ("Przychód (€)", "#,##0 €"),
    ("Śr. wartość koszyka - kupujący (€)", "#,##0.00 €"),
    ("Przychód na użytkownika (€)", "#,##0.00 €"),
    ("Udział w przychodzie", "0.0%"),
]


def formuly(n, kryteria, r):
    """kryteria: lista (litera_kolumny_w_Dane, adres_komorki_z_wartoscia)."""
    warunki = ",".join(f"{zakres(k, n)},{adr}" for k, adr in kryteria)
    zak = f"{zakres('J', n)},\"Tak\""
    return [
        f"=COUNTIFS({warunki})",
        f"=COUNTIFS({warunki},{zak})",
        f"=IFERROR(C{r}/B{r},0)" if len(kryteria) == 1 else f"=IFERROR(D{r}/C{r},0)",
        f"=AVERAGEIFS({zakres('F', n)},{warunki})",
        f"=AVERAGEIFS({zakres('G', n)},{warunki})",
        f"=AVERAGEIFS({zakres('H', n)},{warunki})",
        f"=SUMIFS({zakres('I', n)},{warunki})",
    ]


def arkusz_podsumowanie(wb, df):
    ws = wb.create_sheet("Podsumowanie")
    n = len(df)
    ws["A1"] = "Podsumowanie wskaźników wg filtrów (formuły liczone z arkusza Dane)"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = ("Konwersja = zakupy / użytkownicy; śr. wartość koszyka liczona tylko "
                "dla kupujących; przychód na użytkownika = przychód / wszyscy użytkownicy.")
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18

    def naglowek(r, etykiety):
        for c, tekst in enumerate(etykiety, 1):
            cell = ws.cell(row=r, column=c, value=tekst)
            cell.font = Font(name=FONT, bold=True, color="FFFFFF")
            cell.fill = NAGLOWEK
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = RAMKA
        ws.row_dimensions[r].height = 42

    def formatuj(r, pierwsza_metryka, ile=len(METRYKI)):
        for c in range(1, pierwsza_metryka + ile):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT)
            cell.border = RAMKA
            if c >= pierwsza_metryka:
                cell.number_format = METRYKI[c - pierwsza_metryka][1]

    r = 4
    litery = {"zrodlo_ruchu": "B", "urzadzenie": "C", "kraj": "D", "typ_uzytkownika": "E"}

    # Pojedyncze filtry
    for kol, nazwa in FILTRY:
        ws.cell(row=r, column=1, value=f"Filtr: {nazwa}").font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=1).fill = SEKCJA
        r += 1
        naglowek(r, [nazwa] + [m for m, _ in METRYKI])
        r += 1
        pierwszy = r
        wartosci = list(dict.fromkeys(df[kol]))
        for v in wartosci:
            ws.cell(row=r, column=1, value=v)
            for c, f in enumerate(formuly(n, [(litery[kol], f"$A{r}")], r), 2):
                ws.cell(row=r, column=c, value=f)
            ws.cell(row=r, column=9, value=f"=IFERROR(H{r}/C{r},\"brak zakupów\")")
            ws.cell(row=r, column=10, value=f"=IFERROR(H{r}/B{r},0)")
            ws.cell(row=r, column=11, value=f"=IFERROR(H{r}/SUM($H${pierwszy}:$H${pierwszy + len(wartosci) - 1}),0)")
            formatuj(r, 2)
            r += 1
        # Wiersz "Razem"
        ws.cell(row=r, column=1, value="Razem")
        ws.cell(row=r, column=2, value=f"=SUM(B{pierwszy}:B{r - 1})")
        ws.cell(row=r, column=3, value=f"=SUM(C{pierwszy}:C{r - 1})")
        ws.cell(row=r, column=4, value=f"=IFERROR(C{r}/B{r},0)")
        ws.cell(row=r, column=5, value=f"=AVERAGE({zakres('F', n)})")
        ws.cell(row=r, column=6, value=f"=AVERAGE({zakres('G', n)})")
        ws.cell(row=r, column=7, value=f"=AVERAGE({zakres('H', n)})")
        ws.cell(row=r, column=8, value=f"=SUM(H{pierwszy}:H{r - 1})")
        ws.cell(row=r, column=9, value=f"=IFERROR(H{r}/C{r},0)")
        ws.cell(row=r, column=10, value=f"=IFERROR(H{r}/B{r},0)")
        ws.cell(row=r, column=11, value=f"=IFERROR(H{r}/H{r},0)")
        formatuj(r, 2)
        for c in range(1, 12):
            ws.cell(row=r, column=c).font = Font(name=FONT, bold=True)
        r += 3

    # Filtry krzyżowe: wymiar x Nowy/Powracający
    for kol, nazwa in FILTRY[:3]:
        ws.cell(row=r, column=1, value=f"Filtr: {nazwa} × Nowy/Powracający").font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=1).fill = SEKCJA
        r += 1
        naglowek(r, [nazwa, "Nowy/Powracający"] + [m for m, _ in METRYKI[:-1]])
        r += 1
        for v in dict.fromkeys(df[kol]):
            for typ in ("Nowy", "Powracający"):
                ws.cell(row=r, column=1, value=v)
                ws.cell(row=r, column=2, value=typ)
                kryt = [(litery[kol], f"$A{r}"), ("E", f"$B{r}")]
                f = formuly(n, kryt, r)
                ws.cell(row=r, column=3, value=f[0])
                ws.cell(row=r, column=4, value=f[1])
                ws.cell(row=r, column=5, value=f[2])
                for c, fr in zip((6, 7, 8, 9), f[3:]):
                    ws.cell(row=r, column=c, value=f"=IFERROR({fr[1:]},\"-\")")
                ws.cell(row=r, column=10, value=f"=IFERROR(I{r}/D{r},\"brak zakupów\")")
                ws.cell(row=r, column=11, value=f"=IFERROR(I{r}/C{r},\"-\")")
                formatuj(r, 3, len(METRYKI) - 1)
                r += 1
        r += 2

    for c in range(3, 13):
        ws.column_dimensions[get_column_letter(c)].width = 15
    ws.freeze_panes = "B4"


def zbuduj_xlsx(df):
    wb = Workbook()
    arkusz_dane(wb, df)
    arkusz_podsumowanie(wb, df)
    wb.save(XLSX)
    print(f"\nZapisano {XLSX.name}")


if __name__ == "__main__":
    dane = pd.read_csv(DANE)
    dane["kupil"] = (dane.zakup == "Tak").astype(int)
    wypisz(dane)
    zbuduj_xlsx(dane)
